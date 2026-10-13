#!/usr/bin/env python3
"""Inventaría RFP/ y data-samples/, extrae texto por página y renderiza a imagen
las páginas en las que el texto pierde información (tablas, diagramas, escaneados).

Uso:
  uv run --group se inventory_corpus.py [raíz] [--render auto|all|none] [--dpi 110]
                             [--max-render 200] [--full]
  uv run --group se inventory_corpus.py [raíz] --check     # exit 1 si el manifiesto está obsoleto

Salidas en proposal/_meta/:
  corpus-manifest.json                 inventario, métodos, páginas marcadas y motivos
  corpus/<doc>.txt                     texto con marcas '--- página N ---'
  corpus/pages/<doc>/p-012.png         páginas renderizadas
  corpus/tables/<doc>/p-012-t1.md      tablas extraídas a Markdown
  corpus/converted/<doc>.pdf           DOCX/PPTX/ODT convertidos con LibreOffice

Modo --render:
  auto  (defecto) renderiza páginas con poco texto, tablas, imágenes grandes, dibujo
        vectorial denso o mala codificación; las presentaciones se renderizan enteras.
  all   renderiza todas las páginas (documentos cortos o muy visuales).
  none  solo texto.

Herramientas (opcionales, se degrada con aviso si faltan):
  poppler-utils (pdftotext, pdftoppm, pdfinfo) · pdfplumber · pypdf
  LibreOffice (soffice) para DOCX/PPTX/ODT · python-docx / python-pptx / openpyxl como respaldo
Por defecto es incremental: no reprocesa ficheros cuyo hash y opciones no han cambiado (--full fuerza).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

TEXT_EXT = {".md", ".txt", ".csv", ".tsv", ".json", ".yaml", ".yml", ".xml", ".html", ".htm", ".sql", ".eml"}
OFFICE_TO_PDF = {".docx", ".doc", ".pptx", ".ppt", ".odt", ".odp", ".rtf"}
SLIDES = {".pptx", ".ppt", ".odp"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".tif", ".tiff", ".bmp"}
DATA_EXT = {".csv", ".tsv", ".parquet", ".xlsx", ".xls", ".json", ".jsonl"}

MIN_CHARS = 200            # menos texto → probable escaneado, diagrama o portada
IMAGE_AREA_RATIO = 0.30    # imágenes que ocupan > 30 % de la página
VECTOR_OBJECTS = 80        # líneas/curvas/rectángulos → diagrama dibujado
BAD_CHAR_RATIO = 0.10      # caracteres no imprimibles o de sustitución


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def has(cmd: str) -> bool:
    return shutil.which(cmd) is not None


def soffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def safe_name(rel: str) -> str:
    return re.sub(r"[^\w.-]+", "_", rel.replace("/", "__"))


# ---------------------------------------------------------------- PDF por página

def pdf_page_texts(pdf: Path) -> tuple[list[str], str]:
    if has("pdftotext"):
        r = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", str(pdf), "-"], capture_output=True, timeout=600)
        if r.returncode == 0:
            pages = r.stdout.decode("utf-8", "replace").split("\f")
            if pages and not pages[-1].strip():
                pages = pages[:-1]
            if pages:
                return pages, "pdftotext"
    try:
        import pdfplumber  # type: ignore
        with pdfplumber.open(str(pdf)) as doc:
            return [(p.extract_text() or "") for p in doc.pages], "pdfplumber"
    except ImportError:
        pass
    try:
        from pypdf import PdfReader  # type: ignore
        return [(p.extract_text() or "") for p in PdfReader(str(pdf)).pages], "pypdf"
    except ImportError:
        return [], "sin extractor de PDF (instala poppler-utils o ejecuta /se-init para pdfplumber)"


def bad_ratio(text: str) -> float:
    t = re.sub(r"\s", "", text)
    if not t:
        return 0.0
    bad = sum(1 for c in t if c == "\ufffd" or (not c.isprintable()) or ("\ue000" <= c <= "\uf8ff"))
    return bad / len(t)


def to_md_table(rows: list[list]) -> str:
    rows = [[("" if c is None else str(c)).replace("\n", " ").replace("|", "/").strip() for c in r] for r in rows if r]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out) + "\n"


def analyse_pdf(pdf: Path, texts: list[str], tables_dir: Path, root: Path) -> list[dict]:
    pages = [{"n": i + 1, "chars": len(t.strip()), "reasons": [], "tables": []} for i, t in enumerate(texts)]
    for p, t in zip(pages, texts):
        if p["chars"] < MIN_CHARS:
            p["reasons"].append("poco texto")
        if bad_ratio(t) > BAD_CHAR_RATIO:
            p["reasons"].append("codificación defectuosa")
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        return pages
    with pdfplumber.open(str(pdf)) as doc:
        for i, page in enumerate(doc.pages):
            if i >= len(pages):
                break
            p = pages[i]
            area = float(page.width * page.height) or 1.0
            img_area = sum(abs((im["x1"] - im["x0"]) * (im["bottom"] - im["top"])) for im in page.images)
            img_ratio = img_area / area
            vectors = len(page.lines) + len(page.curves) + len(page.rects)
            if img_ratio > IMAGE_AREA_RATIO:
                p["reasons"].append("imagen o diagrama")
            if "poco texto" in p["reasons"] and p["chars"] > 0 and img_ratio < 0.05 and vectors < 10:
                p["reasons"].remove("poco texto")  # página corta sin elementos visuales: el texto basta
            try:
                found = page.extract_tables()
            except Exception:  # noqa: BLE001
                found = []
            found = [t for t in found if t and len(t) >= 2 and max(len(r) for r in t) >= 2]
            if found:
                p["reasons"].append("tabla")
                tables_dir.mkdir(parents=True, exist_ok=True)
                for k, tbl in enumerate(found, 1):
                    md = to_md_table(tbl)
                    if md:
                        f = tables_dir / f"p-{i + 1:03d}-t{k}.md"
                        f.write_text(md, encoding="utf-8")
                        p["tables"].append(f.relative_to(root).as_posix())
            elif vectors > VECTOR_OBJECTS:
                p["reasons"].append("dibujo vectorial")
    return pages


def render_pages(pdf: Path, numbers: list[int], out_dir: Path, dpi: int, root: Path) -> dict[int, str]:
    if not numbers or not has("pdftoppm"):
        return {}
    out_dir.mkdir(parents=True, exist_ok=True)
    done = {}
    # agrupa en rangos contiguos para minimizar llamadas
    ranges, start = [], numbers[0]
    for a, b in zip(numbers, numbers[1:] + [None]):
        if b != a + 1:
            ranges.append((start, a))
            start = b
    for first, last in ranges:
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(["pdftoppm", "-png", "-r", str(dpi), "-f", str(first), "-l", str(last), str(pdf),
                            str(Path(tmp) / "p")], capture_output=True, timeout=900)
            for f in Path(tmp).glob("p-*.png"):
                n = int(re.search(r"-(\d+)\.png$", f.name)[1])
                target = out_dir / f"p-{n:03d}.png"
                shutil.move(str(f), target)
                done[n] = target.relative_to(root).as_posix()
    return done


def process_pdf(pdf: Path, entry: dict, root: Path, meta: Path, mode: str, dpi: int, budget: list[int],
                force_all: bool = False) -> str:
    texts, method = pdf_page_texts(pdf)
    entry["method"] = method
    if not texts:
        entry["extracted"] = False
        return ""
    stem = safe_name(entry["path"])
    for sub in ("pages", "tables"):
        shutil.rmtree(meta / "corpus" / sub / stem, ignore_errors=True)
    pages = analyse_pdf(pdf, texts, meta / "corpus/tables" / stem, root)
    if mode == "all" or force_all:
        targets = [p["n"] for p in pages]
        for p in pages:
            tag = "presentación" if force_all else "modo all"
            if tag not in p["reasons"]:
                p["reasons"].append(tag)
    elif mode == "auto":
        targets = [p["n"] for p in pages if p["reasons"]]
    else:
        targets = []
    allowed = targets[: max(0, budget[0])]
    budget[0] -= len(allowed)
    rendered = render_pages(pdf, allowed, meta / "corpus/pages" / stem, dpi, root) if mode != "none" else {}
    for p in pages:
        if p["n"] in rendered:
            p["image"] = rendered[p["n"]]
        elif p["n"] in targets:
            p["image"] = None
            p["reasons"].append("sin renderizar (límite --max-render o falta pdftoppm)")
    entry["pages"] = len(pages)
    entry["flagged_pages"] = [p for p in pages if p["reasons"] or p["tables"]]
    entry["rendered"] = len(rendered)
    entry["tables"] = sum(len(p["tables"]) for p in pages)
    scanned = sum(1 for p in pages if p["chars"] < MIN_CHARS)
    if pages and not force_all and scanned / len(pages) > 0.8:
        entry["warning"] = "documento probablemente escaneado: leer desde las imágenes de página"
    return "".join(f"\n--- página {i + 1} ---\n{t}" for i, t in enumerate(texts))


# ---------------------------------------------------------------- otros formatos

def office_to_pdf(src: Path, out_dir: Path) -> Path | None:
    exe = soffice()
    if not exe:
        return None
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run([exe, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
                        "--outdir", str(out_dir), str(src)], capture_output=True, timeout=600)
    pdf = out_dir / (src.stem + ".pdf")
    return pdf if pdf.exists() else None


def fallback_office_text(src: Path) -> tuple[str | None, str]:
    ext = src.suffix.lower()
    try:
        if ext == ".docx":
            import docx  # type: ignore
            d = docx.Document(str(src))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                parts += ["| " + " | ".join(c.text.strip() for c in r.cells) + " |" for r in t.rows]
            return "\n".join(parts), "python-docx (sin renderizado: instala LibreOffice)"
        if ext == ".pptx":
            from pptx import Presentation  # type: ignore
            parts = []
            for i, slide in enumerate(Presentation(str(src)).slides, 1):
                parts.append(f"\n--- diapositiva {i} ---")
                parts += [s.text_frame.text for s in slide.shapes if getattr(s, "has_text_frame", False)]
            return "\n".join(parts), "python-pptx (sin renderizado: instala LibreOffice)"
    except ImportError:
        pass
    return None, "instala LibreOffice (o python-docx / python-pptx como respaldo)"


def xlsx_text(src: Path) -> tuple[str | None, str]:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        return None, "falta openpyxl (ejecuta /se-init)"
    wb = openpyxl.load_workbook(str(src), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"\n--- hoja {ws.title} ---")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 500:
                parts.append("... (truncado a 500 filas; usa profile_dataset.py para datos)")
                break
            parts.append("| " + " | ".join("" if v is None else str(v) for v in row) + " |")
    return "\n".join(parts), "openpyxl"


def process_file(src: Path, entry: dict, root: Path, meta: Path, a, budget: list[int]) -> None:
    ext = src.suffix.lower()
    text, is_sample = None, entry["path"].startswith("data-samples/")
    if is_sample and ext in DATA_EXT:
        entry.update(extracted=False, method="usar profile_dataset.py")
        return
    if ext in TEXT_EXT:
        text, entry["method"] = src.read_text(encoding="utf-8", errors="replace"), "texto plano"
    elif ext == ".pdf":
        text = process_pdf(src, entry, root, meta, a.render, a.dpi, budget)
    elif ext in OFFICE_TO_PDF:
        pdf = office_to_pdf(src, meta / "corpus/converted" / safe_name(str(Path(entry["path"]).parent)))
        if pdf:
            entry["converted_pdf"] = pdf.relative_to(root).as_posix()
            text = process_pdf(pdf, entry, root, meta, a.render, a.dpi, budget, force_all=(ext in SLIDES and a.render != "none"))
            entry["method"] = f"LibreOffice → PDF · {entry.get('method', '')}"
        else:
            text, entry["method"] = fallback_office_text(src)
    elif ext in {".xlsx", ".xlsm"}:
        text, entry["method"] = xlsx_text(src)
    elif ext in IMAGE_EXT:
        entry.update(extracted=False, method="imagen: leer directamente", image=entry["path"])
        return
    else:
        entry["method"] = f"formato {ext or 'sin extensión'} no soportado"
    if text and text.strip():
        out = meta / "corpus" / (safe_name(entry["path"]) + ".txt")
        out.write_text(text, encoding="utf-8")
        entry.update(extracted=True, text_path=out.relative_to(root).as_posix(), chars=len(text))
    else:
        entry.setdefault("extracted", False)


# ---------------------------------------------------------------- main

def scan(root: Path) -> list[dict]:
    items = []
    for base in ("RFP", "data-samples"):
        d = root / base
        if d.is_dir():
            for p in sorted(d.rglob("*")):
                if p.is_file() and not p.name.startswith((".", "~$")):
                    items.append({"path": p.relative_to(root).as_posix(), "size": p.stat().st_size, "sha256": sha256(p)})
    return items


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("root", nargs="?", default=".", help="raíz del workspace (por defecto, el directorio actual)")
    ap.add_argument("--render", choices=["auto", "all", "none"], default="auto")
    ap.add_argument("--dpi", type=int, default=110)
    ap.add_argument("--max-render", type=int, default=200, help="máximo de páginas a renderizar en total")
    ap.add_argument("--check", action="store_true")
    ap.add_argument("--full", action="store_true", help="reprocesa todo aunque no haya cambios")
    a = ap.parse_args()

    root = Path(a.root).resolve()
    if root.name == "RFP":
        root = root.parent
    meta = root / "proposal/_meta"
    manifest_path = meta / "corpus-manifest.json"
    current = scan(root)
    if not current:
        print("ERROR: no hay ficheros en RFP/ ni data-samples/")
        return 1
    old = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else None

    if a.check:
        if not old:
            print("OBSOLETO: no existe corpus-manifest.json")
            return 1
        before = {i["path"]: i["sha256"] for i in old["files"]}
        now = {i["path"]: i["sha256"] for i in current}
        if before != now:
            changed = sorted({k for k in before.keys() | now.keys() if before.get(k) != now.get(k)})
            print("OBSOLETO: cambios en " + ", ".join(changed))
            return 1
        print("OK: manifiesto al día")
        return 0

    opts = {"render": a.render, "dpi": a.dpi}
    reuse = {}
    if old and not a.full and old.get("options") == opts:
        reuse = {f["path"]: f for f in old["files"]}
    if a.full or not reuse:
        for sub in ("corpus",):
            shutil.rmtree(meta / sub, ignore_errors=True)
    (meta / "corpus").mkdir(parents=True, exist_ok=True)

    budget = [a.max_render]
    files, processed = [], 0
    for item in current:
        prev = reuse.get(item["path"])
        if prev and prev["sha256"] == item["sha256"] and (not prev.get("text_path") or (root / prev["text_path"]).exists()):
            files.append(prev)
            budget[0] -= prev.get("rendered", 0)
            continue
        entry = dict(item, kind="muestra-datos" if item["path"].startswith("data-samples/") else "documento")
        try:
            process_file(root / item["path"], entry, root, meta, a, budget)
        except Exception as exc:  # noqa: BLE001
            entry.update(extracted=False, method=f"error: {exc}")
        files.append(entry)
        processed += 1

    tools = {"pdftotext": has("pdftotext"), "pdftoppm": has("pdftoppm"), "libreoffice": bool(soffice())}
    try:
        import pdfplumber  # type: ignore # noqa: F401
        tools["pdfplumber"] = True
    except ImportError:
        tools["pdfplumber"] = False
    manifest = {"generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "root": str(root),
                "options": opts, "tools": tools, "files": files}
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    ok = sum(1 for f in files if f.get("extracted"))
    rendered = sum(f.get("rendered", 0) for f in files)
    tables = sum(f.get("tables", 0) for f in files)
    print(f"manifiesto: {manifest_path}")
    print(f"ficheros: {len(files)} (procesados {processed}, reutilizados {len(files) - processed}) · "
          f"con texto: {ok} · páginas renderizadas: {rendered} · tablas: {tables}")
    missing = [k for k, v in tools.items() if not v]
    if missing:
        print("AVISO: herramientas no disponibles → " + ", ".join(missing) + " (calidad degradada)")
    if budget[0] < 0 or any("límite" in r for f in files for p in f.get("flagged_pages", []) for r in p["reasons"]):
        print("AVISO: se alcanzó --max-render; hay páginas marcadas sin imagen")
    for f in files:
        if f.get("warning"):
            print(f"  ! {f['path']}: {f['warning']}")
        if not f.get("extracted") and f.get("method") != "usar profile_dataset.py":
            print(f"  - {f['path']}: {f.get('method')}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
