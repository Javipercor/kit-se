"""Definiciones de estilos y factory para crearlos."""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


@dataclass(frozen=True)
class Styles:
    """Contiene todos los estilos utilizados en el workbook."""
    thin_border: Border
    center: Alignment
    left: Alignment
    left_no_wrap: Alignment
    right: Alignment

    header_font_white: Font
    small_font: Font
    title_bold: Font

    fill_header_blue: PatternFill
    fill_row_blue: PatternFill
    fill_row_blue_alt: PatternFill
    fill_block_light: PatternFill
    fill_orange: PatternFill
    fill_yellow: PatternFill
    fill_white: PatternFill


class StyleFactory:
    """Factory para crear instancias de Styles."""

    @staticmethod
    def create() -> Styles:
        """Crea y retorna una instancia de Styles con todos los estilos predefinidos."""
        thin_side: Side = Side(style="thin", color="BFBFBF")
        thin_border: Border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        center: Alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left: Alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        left_no_wrap: Alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
        right: Alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

        title_bold: Font = Font(name="Calibri", size=11, bold=True)
        header_font_white: Font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        small_font: Font = Font(name="Calibri", size=9)

        fill_header_blue: PatternFill = PatternFill("solid", fgColor="2F5597")
        fill_row_blue: PatternFill = PatternFill("solid", fgColor="5B9BD5")
        fill_row_blue_alt: PatternFill = PatternFill("solid", fgColor="4472C4")
        fill_block_light: PatternFill = PatternFill("solid", fgColor="D9E1F2")
        fill_orange: PatternFill = PatternFill("solid", fgColor="F4B183")
        fill_yellow: PatternFill = PatternFill("solid", fgColor="FFFF00")
        fill_white: PatternFill = PatternFill("solid", fgColor="FFFFFF")

        return Styles(
            thin_border=thin_border,
            center=center,
            left=left,
            left_no_wrap=left_no_wrap,
            right=right,
            header_font_white=header_font_white,
            small_font=small_font,
            title_bold=title_bold,
            fill_header_blue=fill_header_blue,
            fill_row_blue=fill_row_blue,
            fill_row_blue_alt=fill_row_blue_alt,
            fill_block_light=fill_block_light,
            fill_orange=fill_orange,
            fill_yellow=fill_yellow,
            fill_white=fill_white,
        )
