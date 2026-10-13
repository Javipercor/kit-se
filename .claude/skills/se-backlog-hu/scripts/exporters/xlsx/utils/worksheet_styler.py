"""Utilidades para aplicar estilos a celdas y rangos."""

from __future__ import annotations

from typing import Optional

from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from ..styles import Styles


class WorksheetStyler:
    """Helper para aplicar estilos a celdas y rangos en un worksheet."""

    def __init__(self, ws: Worksheet, styles: Styles) -> None:
        self._ws: Worksheet = ws
        self._styles: Styles = styles

    def set_cell(
        self,
        address: str,
        value: object | None = None,
        *,
        font: Optional[Font] = None,
        fill: Optional[PatternFill] = None,
        alignment: Optional[Alignment] = None,
        border: Optional[Border] = None,
        number_format: Optional[str] = None,
    ) -> None:
        """Establece valor y estilos en una celda."""
        cell = self._ws[address]
        if value is not None:
            cell.value = value
        if font is not None:
            cell.font = font
        if fill is not None:
            cell.fill = fill
        if alignment is not None:
            cell.alignment = alignment
        if border is not None:
            cell.border = border
        if number_format is not None:
            cell.number_format = number_format

    def apply_range(
        self,
        range_ref: str,
        *,
        font: Optional[Font] = None,
        fill: Optional[PatternFill] = None,
        alignment: Optional[Alignment] = None,
        border: Optional[Border] = None,
    ) -> None:
        """Aplica estilos a un rango de celdas."""
        for row in self._ws[range_ref]:
            for cell in row:
                if font is not None:
                    cell.font = font
                if fill is not None:
                    cell.fill = fill
                if alignment is not None:
                    cell.alignment = alignment
                if border is not None:
                    cell.border = border

    def merge_and_style(
        self,
        range_ref: str,
        *,
        value: object | None = None,
        font: Optional[Font] = None,
        fill: Optional[PatternFill] = None,
        alignment: Optional[Alignment] = None,
        border: Optional[Border] = None,
    ) -> None:
        """Fusiona celdas y aplica estilos."""
        self._ws.merge_cells(range_ref)
        top_left = range_ref.split(":")[0]
        self.set_cell(top_left, value=value, font=font, fill=fill, alignment=alignment, border=border)
        self.apply_range(range_ref, font=font, fill=fill, alignment=alignment, border=border)
